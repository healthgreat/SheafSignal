#!/usr/bin/env python
"""Build a journal-upload package for the SheafSignal manuscript.

Author: SheafSignal contributors
Date: 2026-05-01
Purpose: generate DOCX upload artifacts, a manifest, and a preflight checklist
from auditable Markdown/TSV manuscript sources.
"""

from __future__ import annotations

import argparse
import hashlib
import shutil
from pathlib import Path
import re
import subprocess
import tempfile

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUTPUT_FILES = {
    "main_docx": "SheafSignal_main_manuscript_v2.docx",
    "cover_docx": "SheafSignal_cover_letter_NatureMethods.docx",
    "supplement_docx": "SheafSignal_supplementary_information.docx",
    "manifest": "submission_upload_manifest.tsv",
    "preflight": "submission_upload_preflight_checklist.tsv",
    "readme": "README.md",
    "text_check": "DOCX_TEXT_EXTRACTION_CHECK.md",
    "layout_check": "DOCX_LAYOUT_CHECK.tsv",
    "figure_manifest": "main_figure_upload_manifest.tsv",
}


EXPECTED_MAIN_STRINGS = (
    "SheafSignal maps communication frustration",
    "not a full pretrained NicheNet network benchmark",
    "does not claim clinical",
)


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_text_atomic(path: Path, text: str) -> None:
    tmp_path = path.with_suffix(path.suffix + ".tmp")
    tmp_path.write_text(text, encoding="utf-8")
    tmp_path.replace(path)


def _write_table_atomic(path: Path, table: pd.DataFrame) -> None:
    tmp_path = path.with_suffix(path.suffix + ".tmp")
    table.to_csv(tmp_path, sep="\t", index=False)
    tmp_path.replace(path)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8") if path.exists() else ""


def _read_tsv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    return pd.read_csv(path, sep="\t")


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = document.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def _split_markdown_table_row(line: str) -> list[str]:
    stripped = line.strip().strip("|")
    return [cell.strip() for cell in stripped.split("|")]


def _is_markdown_table_separator(line: str) -> bool:
    cells = _split_markdown_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [_split_markdown_table_row(line) for line in lines]
    if len(rows) < 2:
        return
    headers = rows[0]
    data = (
        rows[2:]
        if len(rows) > 1 and _is_markdown_table_separator(lines[1])
        else rows[1:]
    )
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for idx, header in enumerate(headers):
        run = table.rows[0].cells[idx].paragraphs[0].add_run(header)
        run.bold = True
    for row in data:
        cells = table.add_row().cells
        for idx, value in enumerate(row[: len(headers)]):
            cells[idx].text = value


def add_markdown(document: Document, markdown_text: str) -> None:
    lines = markdown_text.splitlines()
    in_code = False
    code_buffer: list[str] = []
    i = 0
    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                para = document.add_paragraph()
                run = para.add_run("\n".join(code_buffer))
                run.font.name = "Courier New"
                run.font.size = Pt(9)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if (
            stripped.startswith("|")
            and i + 1 < len(lines)
            and _is_markdown_table_separator(lines[i + 1])
        ):
            table_lines = [stripped, lines[i + 1].strip()]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            _add_markdown_table(document, table_lines)
            continue

        if stripped.startswith("# "):
            para = document.add_heading(stripped[2:].strip(), level=1)
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif stripped.startswith("## "):
            document.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith("### "):
            document.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith("- "):
            document.add_paragraph(stripped[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\.\s+", stripped):
            document.add_paragraph(
                re.sub(r"^\d+\.\s+", "", stripped), style="List Number"
            )
        else:
            document.add_paragraph(stripped)
        i += 1


def build_docx_from_markdown(markdown_path: Path, output_path: Path) -> Path:
    document = Document()
    _configure_document(document)
    text = _read_text(markdown_path)
    add_markdown(document, text)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = output_path.with_suffix(output_path.suffix + ".tmp")
    document.save(tmp_path)
    tmp_path.replace(output_path)
    return output_path


def extract_docx_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _copy_atomic(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = destination.with_suffix(destination.suffix + ".tmp")
    shutil.copy2(source, tmp_path)
    tmp_path.replace(destination)


def build_main_figure_upload_manifest(
    *,
    root: Path,
    figure_manifest_path: Path,
    output_dir: Path,
) -> pd.DataFrame:
    if not figure_manifest_path.exists():
        return pd.DataFrame(
            [
                {
                    "figure_id": "NA",
                    "source_path": figure_manifest_path.relative_to(root).as_posix(),
                    "upload_path": "NA",
                    "status": "missing_figure_manifest",
                    "sha256": "NA",
                    "size_mb": "NA",
                }
            ]
        )
    manifest = pd.read_csv(figure_manifest_path, sep="\t")
    if (
        manifest.empty
        or "figure_id" not in manifest.columns
        or "path" not in manifest.columns
    ):
        return pd.DataFrame(
            [
                {
                    "figure_id": "NA",
                    "source_path": figure_manifest_path.relative_to(root).as_posix(),
                    "upload_path": "NA",
                    "status": "invalid_figure_manifest",
                    "sha256": "NA",
                    "size_mb": "NA",
                }
            ]
        )
    rows = []
    figure_dir = output_dir / "main_figures"
    for row in manifest.to_dict(orient="records"):
        figure_id = str(row.get("figure_id", ""))
        if not figure_id.startswith("fig"):
            continue
        rel_source = Path(str(row.get("path", "")))
        source = root / rel_source
        destination = figure_dir / source.name
        if not source.exists():
            rows.append(
                {
                    "figure_id": figure_id,
                    "source_path": rel_source.as_posix(),
                    "upload_path": destination.relative_to(root).as_posix(),
                    "status": "missing_source_figure",
                    "sha256": "NA",
                    "size_mb": "NA",
                }
            )
            continue
        _copy_atomic(source, destination)
        rows.append(
            {
                "figure_id": figure_id,
                "source_path": rel_source.as_posix(),
                "upload_path": destination.relative_to(root).as_posix(),
                "status": "copied",
                "sha256": sha256_file(destination),
                "size_mb": round(destination.stat().st_size / 1024 / 1024, 6),
            }
        )
    return pd.DataFrame(rows)


def build_manifest(
    *,
    root: Path,
    source_paths: dict[str, Path],
    output_paths: dict[str, Path],
    figure_upload_manifest: pd.DataFrame | None = None,
) -> pd.DataFrame:
    rows = [
        {
            "upload_item": "main_manuscript",
            "file_path": output_paths["main_docx"].relative_to(root).as_posix(),
            "source_path": source_paths["main_md"].as_posix(),
            "submission_role": "Main manuscript DOCX generated from polished v2 Markdown",
            "status": "generated",
        },
        {
            "upload_item": "cover_letter",
            "file_path": output_paths["cover_docx"].relative_to(root).as_posix(),
            "source_path": source_paths["cover_md"].as_posix(),
            "submission_role": "Nature Methods cover letter DOCX",
            "status": "generated",
        },
        {
            "upload_item": "supplementary_information",
            "file_path": output_paths["supplement_docx"].relative_to(root).as_posix(),
            "source_path": source_paths["supplement_md"].as_posix(),
            "submission_role": "Supplementary information DOCX",
            "status": "generated",
        },
        {
            "upload_item": "zenodo_archive",
            "file_path": "release/archives/sheafsignal_zenodo_upload.zip",
            "source_path": "release/zenodo_upload_manifest.tsv",
            "submission_role": "Upload to Zenodo before journal submission",
            "status": "pending_external_deposition",
        },
        {
            "upload_item": "github_archive",
            "file_path": "release/archives/sheafsignal_github_release.zip",
            "source_path": "release/github_release_manifest.tsv",
            "submission_role": "Public GitHub release source archive",
            "status": "pending_external_public_release",
        },
    ]
    if figure_upload_manifest is not None and not figure_upload_manifest.empty:
        for row in figure_upload_manifest.to_dict(orient="records"):
            rows.append(
                {
                    "upload_item": str(row["figure_id"]),
                    "file_path": str(row["upload_path"]),
                    "source_path": str(row["source_path"]),
                    "submission_role": "Main figure PDF upload",
                    "status": str(row["status"]),
                }
            )
    for row in rows:
        path = root / row["file_path"]
        row["exists"] = path.exists()
        row["sha256"] = sha256_file(path) if path.exists() and path.is_file() else "NA"
        row["size_mb"] = (
            round(path.stat().st_size / 1024 / 1024, 6) if path.exists() else "NA"
        )
    return pd.DataFrame(rows)


def _blocker_status(final_blockers: pd.DataFrame, blocker_id: str) -> str:
    if final_blockers.empty:
        return "missing_final_blocker_table"
    rows = final_blockers.loc[final_blockers["blocker_id"].astype(str) == blocker_id]
    if rows.empty:
        return "not_found"
    return str(rows.iloc[0]["status"])


def build_preflight_checklist(
    *,
    root: Path | None = None,
    output_paths: dict[str, Path],
    final_blockers: pd.DataFrame,
    editorial_audit: pd.DataFrame,
    claim_safety_report: str,
    format_audit_report: str,
    text_check_passed: bool,
    layout_check: pd.DataFrame | None = None,
    figure_upload_manifest: pd.DataFrame | None = None,
    figure_quality_audit: pd.DataFrame | None = None,
    supplementary_artifact_audit: pd.DataFrame | None = None,
    benchmark_contract_audit: pd.DataFrame | None = None,
    submission_provenance_audit: pd.DataFrame | None = None,
    method_reporting_audit: pd.DataFrame | None = None,
) -> pd.DataFrame:
    def rel(path: Path) -> str:
        if root is None:
            return path.as_posix()
        try:
            return path.relative_to(root).as_posix()
        except ValueError:
            return path.as_posix()

    editorial_failures = (
        editorial_audit.loc[editorial_audit["status"].astype(str) == "fail"]
        if not editorial_audit.empty and "status" in editorial_audit
        else pd.DataFrame()
    )
    rows = [
        {
            "check_id": "main_manuscript_docx_exists",
            "status": "pass" if output_paths["main_docx"].exists() else "fail",
            "evidence": rel(output_paths["main_docx"]),
            "required_action": "Regenerate upload package.",
        },
        {
            "check_id": "cover_letter_docx_exists",
            "status": "pass" if output_paths["cover_docx"].exists() else "fail",
            "evidence": rel(output_paths["cover_docx"]),
            "required_action": "Regenerate upload package.",
        },
        {
            "check_id": "supplementary_docx_exists",
            "status": "pass" if output_paths["supplement_docx"].exists() else "fail",
            "evidence": rel(output_paths["supplement_docx"]),
            "required_action": "Regenerate upload package.",
        },
        {
            "check_id": "docx_text_extraction",
            "status": "pass" if text_check_passed else "fail",
            "evidence": "DOCX text extraction contains expected boundary strings.",
            "required_action": "Open DOCX manually and inspect formatting if this fails.",
        },
        {
            "check_id": "main_figure_upload_files",
            "status": (
                "pass"
                if figure_upload_manifest is not None
                and not figure_upload_manifest.empty
                and len(figure_upload_manifest) >= 5
                and (figure_upload_manifest["status"].astype(str) == "copied").all()
                else "fail"
            ),
            "evidence": "manuscript/submission_upload_package/main_figure_upload_manifest.tsv",
            "required_action": "Run make_publication_figures.py and rebuild the upload package.",
        },
        {
            "check_id": "main_figure_quality_audit",
            "status": (
                "pass"
                if figure_quality_audit is not None
                and not figure_quality_audit.empty
                and "status" in figure_quality_audit
                and (figure_quality_audit["status"].astype(str) == "pass").all()
                else "pending"
            ),
            "evidence": "manuscript/figure_quality/MAIN_FIGURE_QUALITY_AUDIT.tsv",
            "required_action": "Run check_main_figure_quality.py before upload.",
        },
        {
            "check_id": "supplementary_artifact_audit",
            "status": (
                "pass"
                if supplementary_artifact_audit is not None
                and not supplementary_artifact_audit.empty
                and "status" in supplementary_artifact_audit
                and (supplementary_artifact_audit["status"].astype(str) == "pass").all()
                else "pending"
            ),
            "evidence": "manuscript/supplementary_quality/SUPPLEMENTARY_ARTIFACT_AUDIT.tsv",
            "required_action": "Run check_supplementary_artifacts.py before upload.",
        },
        {
            "check_id": "method_reporting_audit",
            "status": (
                "pass"
                if method_reporting_audit is not None
                and not method_reporting_audit.empty
                and "status" in method_reporting_audit
                and not (method_reporting_audit["status"].astype(str) == "fail").any()
                else "pending"
            ),
            "evidence": "manuscript/method_reporting/METHOD_REPORTING_AUDIT.tsv",
            "required_action": "Run check_method_reporting_readiness.py before upload.",
        },
        {
            "check_id": "benchmark_result_contract_audit",
            "status": (
                "pass"
                if benchmark_contract_audit is not None
                and not benchmark_contract_audit.empty
                and "status" in benchmark_contract_audit
                and not (benchmark_contract_audit["status"].astype(str) == "fail").any()
                else "pending"
            ),
            "evidence": "benchmarks/results/BENCHMARK_RESULT_CONTRACT_AUDIT.tsv",
            "required_action": "Run check_benchmark_result_contracts.py before upload.",
        },
        {
            "check_id": "submission_provenance_audit",
            "status": (
                "pass"
                if submission_provenance_audit is not None
                and not submission_provenance_audit.empty
                and "status" in submission_provenance_audit
                and not (submission_provenance_audit["status"].astype(str) == "fail").any()
                else "pending"
            ),
            "evidence": "manuscript/provenance/SUBMISSION_PROVENANCE_AUDIT.tsv",
            "required_action": "Run check_submission_provenance.py before upload.",
        },
        {
            "check_id": "claim_safety",
            "status": (
                "pass"
                if "Blocking positive claims: 0" in claim_safety_report
                else "fail"
            ),
            "evidence": "manuscript/CLAIM_SAFETY_AUDIT_REPORT.md",
            "required_action": "Resolve blocking positive claims.",
        },
        {
            "check_id": "nature_methods_format",
            "status": (
                "pass" if "FORMAT_LOCALLY_READY" in format_audit_report else "pending"
            ),
            "evidence": "manuscript/NATURE_METHODS_FORMAT_AUDIT_REPORT.md",
            "required_action": "Resolve local format issues and check current journal instructions.",
        },
        {
            "check_id": "v2_editorial_audit",
            "status": "pass" if editorial_failures.empty else "fail",
            "evidence": "manuscript/SCI_MANUSCRIPT_V2_EDITORIAL_AUDIT.tsv",
            "required_action": "Resolve v2 editorial audit failures.",
        },
        {
            "check_id": "zenodo_doi",
            "status": _blocker_status(final_blockers, "zenodo_doi::dataset_manifest"),
            "evidence": "metadata/datasets.tsv",
            "required_action": "Mint Zenodo DOI and run scripts/finalize_zenodo_doi.py.",
        },
        {
            "check_id": "github_public_release",
            "status": _blocker_status(
                final_blockers, "checklist::GitHub repository public release"
            ),
            "evidence": "release/archives/sheafsignal_github_release.zip",
            "required_action": "Publish public GitHub repository or attach release archive.",
        },
        {
            "check_id": "author_metadata",
            "status": _blocker_status(final_blockers, "author_metadata::authors"),
            "evidence": "manuscript/submission_metadata/",
            "required_action": "Fill author names, affiliations, CRediT roles and conflicts.",
        },
    ]
    if layout_check is not None and not layout_check.empty:
        if (layout_check["status"].astype(str) == "pass").all():
            render_status = "pass"
            render_evidence = (
                "manuscript/submission_upload_package/DOCX_LAYOUT_CHECK.tsv"
            )
            render_action = (
                "Perform final manual Word/LibreOffice visual review before upload."
            )
        elif layout_check["status"].astype(str).str.contains("unavailable").all():
            render_status = "visual_render_not_available"
            render_evidence = "DOCX layout render unavailable."
            render_action = (
                "Manual Word/LibreOffice layout review is required before upload."
            )
        else:
            render_status = "fail"
            render_evidence = (
                "manuscript/submission_upload_package/DOCX_LAYOUT_CHECK.tsv"
            )
            render_action = "Inspect DOCX/PDF conversion failures before upload."
    else:
        render_status = "visual_render_not_available"
        render_evidence = "Requires soffice and PyMuPDF for automated visual review."
        render_action = (
            "Manual Word/LibreOffice layout review is required before upload."
        )
    rows.append(
        {
            "check_id": "visual_layout_review",
            "status": render_status,
            "evidence": render_evidence,
            "required_action": render_action,
        }
    )
    return pd.DataFrame(rows)


def build_text_check(output_paths: dict[str, Path]) -> tuple[str, bool]:
    extracted = {
        key: extract_docx_text(path)
        for key, path in output_paths.items()
        if path.suffix == ".docx"
    }
    main_text = extracted.get("main_docx", "")
    expected_results = {
        expected: expected in main_text for expected in EXPECTED_MAIN_STRINGS
    }
    passed = all(expected_results.values())
    lines = [
        "# DOCX Text Extraction Check",
        "",
        f"Decision: `{'PASS' if passed else 'FAIL'}`",
        "",
        "## Expected Main-Manuscript Strings",
        "",
    ]
    for expected, found in expected_results.items():
        lines.append(f"- `{expected}`: {'found' if found else 'missing'}")
    lines.extend(
        [
            "",
            "## Word Counts",
            "",
        ]
    )
    for key, text in extracted.items():
        lines.append(f"- `{key}`: {len(text.split())} extracted words")
    lines.extend(
        [
            "",
            "## Boundary",
            "",
            "This is a text-extraction check, not a full visual layout review.",
            "A manual Word/LibreOffice inspection is still required before upload.",
            "",
        ]
    )
    return "\n".join(lines), passed


def _nonwhite_fraction(samples: bytes, n_channels: int) -> float:
    if not samples or n_channels <= 0:
        return 0.0
    n_pixels = len(samples) // n_channels
    if n_pixels == 0:
        return 0.0
    nonwhite = 0
    for idx in range(0, len(samples), n_channels):
        pixel = samples[idx : idx + n_channels]
        if any(channel < 245 for channel in pixel[:3]):
            nonwhite += 1
    return nonwhite / n_pixels


def _format_rel(path: Path, root: Path | None) -> str:
    if root is None:
        return path.as_posix()
    try:
        return path.relative_to(root).as_posix()
    except ValueError:
        return path.as_posix()


def _layout_unavailable_rows(
    output_paths: dict[str, Path],
    reason: str,
    root: Path | None = None,
) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "docx_key": key,
                "docx_path": _format_rel(path, root),
                "status": "visual_render_unavailable",
                "pdf_path": "NA",
                "n_pages": "NA",
                "first_page_text_chars": "NA",
                "first_page_nonwhite_fraction": "NA",
                "page_width_pt": "NA",
                "page_height_pt": "NA",
                "notes": reason,
            }
            for key, path in output_paths.items()
            if path.suffix == ".docx"
        ]
    )


def build_layout_check(
    output_paths: dict[str, Path],
    root: Path | None = None,
) -> pd.DataFrame:
    soffice = shutil.which("soffice")
    if soffice is None:
        return _layout_unavailable_rows(output_paths, "soffice not found", root=root)
    try:
        import fitz
    except ImportError:
        return _layout_unavailable_rows(
            output_paths,
            "PyMuPDF fitz not found",
            root=root,
        )

    rows: list[dict[str, object]] = []
    with tempfile.TemporaryDirectory(prefix="sheafsignal_docx_layout_") as tmp_dir:
        pdf_dir = Path(tmp_dir)
        for key, path in output_paths.items():
            if path.suffix != ".docx":
                continue
            result = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(pdf_dir),
                    str(path),
                ],
                check=False,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                timeout=120,
            )
            pdf_path = pdf_dir / f"{path.stem}.pdf"
            if result.returncode != 0 or not pdf_path.exists():
                rows.append(
                    {
                        "docx_key": key,
                        "docx_path": _format_rel(path, root),
                        "status": "fail",
                        "pdf_path": "NA",
                        "n_pages": "NA",
                        "first_page_text_chars": "NA",
                        "first_page_nonwhite_fraction": "NA",
                        "page_width_pt": "NA",
                        "page_height_pt": "NA",
                        "notes": (result.stderr or result.stdout).strip()[:500],
                    }
                )
                continue
            pdf = fitz.open(pdf_path)
            n_pages = pdf.page_count
            first_page = pdf[0] if n_pages else None
            if first_page is None:
                text_chars = 0
                nonwhite = 0.0
                width = "NA"
                height = "NA"
            else:
                text_chars = len(first_page.get_text())
                pix = first_page.get_pixmap(
                    matrix=fitz.Matrix(0.25, 0.25),
                    alpha=False,
                )
                nonwhite = _nonwhite_fraction(pix.samples, pix.n)
                width = round(float(first_page.rect.width), 3)
                height = round(float(first_page.rect.height), 3)
            status = (
                "pass"
                if n_pages > 0 and text_chars > 50 and nonwhite > 0.005
                else "fail"
            )
            rows.append(
                {
                    "docx_key": key,
                    "docx_path": _format_rel(path, root),
                    "status": status,
                    "pdf_path": "temporary_pdf_rendered_and_removed",
                    "n_pages": n_pages,
                    "first_page_text_chars": text_chars,
                    "first_page_nonwhite_fraction": round(nonwhite, 6),
                    "page_width_pt": width,
                    "page_height_pt": height,
                    "notes": "soffice_pdf_export_plus_pymupdf_render",
                }
            )
            pdf.close()
    return pd.DataFrame(rows)


def build_readme(preflight: pd.DataFrame) -> str:
    blocking = preflight.loc[
        preflight["status"].astype(str).isin(["fail", "pending", "pending"])
    ]
    return f"""# SheafSignal Submission Upload Package

This folder contains local DOCX upload artifacts generated from auditable
Markdown/TSV sources. The DOCX files are convenience upload products; the
authoritative sources remain the Markdown manuscript, claim tracker, benchmark
tables and release manifests.

## Files

- `SheafSignal_main_manuscript_v2.docx`
- `SheafSignal_cover_letter_NatureMethods.docx`
- `SheafSignal_supplementary_information.docx`
- `submission_upload_manifest.tsv`
- `submission_upload_preflight_checklist.tsv`
- `main_figure_upload_manifest.tsv`
- `DOCX_TEXT_EXTRACTION_CHECK.md`
- `DOCX_LAYOUT_CHECK.tsv`
- `manuscript/figure_quality/MAIN_FIGURE_QUALITY_AUDIT.tsv` remains the
  source quality gate for the copied main-figure PDFs.
- `manuscript/supplementary_quality/SUPPLEMENTARY_ARTIFACT_AUDIT.tsv` remains
  the readability gate for supplementary figures, tables, and support files.
- `benchmarks/results/BENCHMARK_RESULT_CONTRACT_AUDIT.tsv` remains the
  result-table contract gate for benchmark summaries and comparator links.
- `manuscript/provenance/SUBMISSION_PROVENANCE_AUDIT.tsv` remains the
  provenance gate for submission-facing artifacts.
- `manuscript/method_reporting/METHOD_REPORTING_AUDIT.tsv` remains the
  method-reporting gate for reviewer-facing manuscript claims.

## Current Boundary

This package does not make the manuscript submission-ready by itself. Zenodo
DOI, public GitHub URL, author metadata, affiliations, CRediT roles, competing
interests and final journal-system checks remain required.

## Preflight Snapshot

- Rows requiring attention: {len(blocking)}
- Final go/no-go source: `manuscript/NATURE_METHODS_GO_NO_GO_REPORT.md`
"""


def build_package(
    *,
    root: Path,
    output_dir: Path,
    main_md: Path,
    cover_md: Path,
    supplement_md: Path,
) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = {key: output_dir / name for key, name in OUTPUT_FILES.items()}

    build_docx_from_markdown(main_md, paths["main_docx"])
    build_docx_from_markdown(cover_md, paths["cover_docx"])
    build_docx_from_markdown(supplement_md, paths["supplement_docx"])

    output_docx = {
        "main_docx": paths["main_docx"],
        "cover_docx": paths["cover_docx"],
        "supplement_docx": paths["supplement_docx"],
    }
    text_check, text_check_passed = build_text_check(output_docx)
    _write_text_atomic(paths["text_check"], text_check)
    layout_check = build_layout_check(output_docx, root=root)
    _write_table_atomic(paths["layout_check"], layout_check)
    figure_upload_manifest = build_main_figure_upload_manifest(
        root=root,
        figure_manifest_path=root / "manuscript/figure_manifest.tsv",
        output_dir=output_dir,
    )
    _write_table_atomic(paths["figure_manifest"], figure_upload_manifest)

    manifest = build_manifest(
        root=root,
        source_paths={
            "main_md": main_md.relative_to(root),
            "cover_md": cover_md.relative_to(root),
            "supplement_md": supplement_md.relative_to(root),
        },
        output_paths=paths,
        figure_upload_manifest=figure_upload_manifest,
    )
    _write_table_atomic(paths["manifest"], manifest)

    preflight = build_preflight_checklist(
        root=root,
        output_paths=output_docx,
        final_blockers=_read_tsv(root / "manuscript/FINAL_SUBMISSION_BLOCKERS.tsv"),
        editorial_audit=_read_tsv(
            root / "manuscript/SCI_MANUSCRIPT_V2_EDITORIAL_AUDIT.tsv"
        ),
        claim_safety_report=_read_text(
            root / "manuscript/CLAIM_SAFETY_AUDIT_REPORT.md"
        ),
        format_audit_report=_read_text(
            root / "manuscript/NATURE_METHODS_FORMAT_AUDIT_REPORT.md"
        ),
        text_check_passed=text_check_passed,
        layout_check=layout_check,
        figure_upload_manifest=figure_upload_manifest,
        figure_quality_audit=_read_tsv(
            root / "manuscript/figure_quality/MAIN_FIGURE_QUALITY_AUDIT.tsv"
        ),
        supplementary_artifact_audit=_read_tsv(
            root
            / "manuscript/supplementary_quality/SUPPLEMENTARY_ARTIFACT_AUDIT.tsv"
        ),
        benchmark_contract_audit=_read_tsv(
            root / "benchmarks/results/BENCHMARK_RESULT_CONTRACT_AUDIT.tsv"
        ),
        submission_provenance_audit=_read_tsv(
            root / "manuscript/provenance/SUBMISSION_PROVENANCE_AUDIT.tsv"
        ),
        method_reporting_audit=_read_tsv(
            root / "manuscript/method_reporting/METHOD_REPORTING_AUDIT.tsv"
        ),
    )
    _write_table_atomic(paths["preflight"], preflight)
    _write_text_atomic(paths["readme"], build_readme(preflight))
    return paths


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", default=".")
    parser.add_argument("--output-dir", default="manuscript/submission_upload_package")
    parser.add_argument("--main-md", default="manuscript/SCI_MANUSCRIPT_V2_POLISHED.md")
    parser.add_argument(
        "--cover-md", default="manuscript/SCI_COVER_LETTER_NatureMethods.md"
    )
    parser.add_argument(
        "--supplement-md",
        default="manuscript/nature_methods_package/10_supplementary_information_draft.md",
    )
    args = parser.parse_args(argv)

    root = Path(args.root).resolve()
    output_dir = Path(args.output_dir)
    if not output_dir.is_absolute():
        output_dir = root / output_dir
    paths = build_package(
        root=root,
        output_dir=output_dir,
        main_md=root / args.main_md,
        cover_md=root / args.cover_md,
        supplement_md=root / args.supplement_md,
    )
    for path in paths.values():
        print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
