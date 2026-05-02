from pathlib import Path
import importlib.util

from docx import Document
import pandas as pd


def _load_script():
    root = Path(__file__).resolve().parents[1]
    path = root / "scripts" / "build_submission_upload_package.py"
    spec = importlib.util.spec_from_file_location(
        "build_submission_upload_package", path
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_markdown_to_docx_preserves_headings_lists_and_tables(tmp_path):
    script = _load_script()
    md = tmp_path / "draft.md"
    md.write_text(
        "# Title\n\n"
        "Body paragraph.\n\n"
        "## Results\n\n"
        "- Bullet item\n"
        "1. Numbered item\n\n"
        "| A | B |\n"
        "|---|---|\n"
        "| x | y |\n",
        encoding="utf-8",
    )
    output = tmp_path / "draft.docx"

    script.build_docx_from_markdown(md, output)

    assert output.exists()
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Title" in text
    assert "Body paragraph." in text
    assert document.tables[0].cell(1, 0).text == "x"


def test_upload_preflight_marks_external_blockers_as_pending(tmp_path):
    script = _load_script()
    output_paths = {
        "main_docx": tmp_path / "main.docx",
        "cover_docx": tmp_path / "cover.docx",
        "supplement_docx": tmp_path / "supp.docx",
    }
    for path in output_paths.values():
        Document().save(path)
    final_blockers = pd.DataFrame(
        [
            {
                "blocker_id": "zenodo_doi::dataset_manifest",
                "status": "pending",
            },
            {
                "blocker_id": "checklist::GitHub repository public release",
                "status": "pending",
            },
            {
                "blocker_id": "author_metadata::authors",
                "status": "tbd_by_authors",
            },
        ]
    )
    editorial = pd.DataFrame(
        [{"audit_id": "forbidden_positive_claims_absent", "status": "pass"}]
    )

    checklist = script.build_preflight_checklist(
        output_paths=output_paths,
        final_blockers=final_blockers,
        editorial_audit=editorial,
        claim_safety_report="Blocking positive claims: 0",
        format_audit_report="FORMAT_LOCALLY_READY",
        text_check_passed=True,
        layout_check=pd.DataFrame(
            [
                {"docx_key": "main_docx", "status": "pass"},
                {"docx_key": "cover_docx", "status": "pass"},
                {"docx_key": "supplement_docx", "status": "pass"},
            ]
        ),
        figure_upload_manifest=pd.DataFrame(
            [{"figure_id": f"fig{idx}", "status": "copied"} for idx in range(1, 6)]
        ),
    )

    statuses = dict(zip(checklist["check_id"], checklist["status"]))
    assert statuses["docx_text_extraction"] == "pass"
    assert statuses["visual_layout_review"] == "pass"
    assert statuses["main_figure_upload_files"] == "pass"
    assert statuses["main_figure_quality_audit"] == "pending"
    assert statuses["supplementary_artifact_audit"] == "pending"
    assert statuses["method_reporting_audit"] == "pending"
    assert statuses["benchmark_result_contract_audit"] == "pending"
    assert statuses["submission_provenance_audit"] == "pending"
    assert statuses["zenodo_doi"] == "pending"
    assert statuses["github_public_release"] == "pending"
    assert statuses["author_metadata"] == "tbd_by_authors"


def test_submission_upload_package_writes_docx_manifest_and_checks(tmp_path):
    script = _load_script()
    root = tmp_path
    manuscript = root / "manuscript"
    manuscript.mkdir()
    main_md = manuscript / "SCI_MANUSCRIPT_V2_POLISHED.md"
    cover_md = manuscript / "SCI_COVER_LETTER_NatureMethods.md"
    supplement_md = manuscript / "supplement.md"
    main_md.write_text(
        "# SheafSignal maps communication frustration\n\n"
        "This is not a full pretrained NicheNet network benchmark.\n"
        "This manuscript does not claim clinical utility.\n",
        encoding="utf-8",
    )
    cover_md.write_text("# Cover Letter\n\nDear Editors,\n", encoding="utf-8")
    supplement_md.write_text("# Supplement\n\nAdditional methods.\n", encoding="utf-8")
    pd.DataFrame(
        [
            {
                "blocker_id": "zenodo_doi::dataset_manifest",
                "status": "pending",
            },
            {
                "blocker_id": "checklist::GitHub repository public release",
                "status": "pending",
            },
            {
                "blocker_id": "author_metadata::authors",
                "status": "tbd_by_authors",
            },
        ]
    ).to_csv(manuscript / "FINAL_SUBMISSION_BLOCKERS.tsv", sep="\t", index=False)
    pd.DataFrame(
        [{"audit_id": "forbidden_positive_claims_absent", "status": "pass"}]
    ).to_csv(
        manuscript / "SCI_MANUSCRIPT_V2_EDITORIAL_AUDIT.tsv", sep="\t", index=False
    )
    (manuscript / "CLAIM_SAFETY_AUDIT_REPORT.md").write_text(
        "Blocking positive claims: 0\n",
        encoding="utf-8",
    )
    (manuscript / "NATURE_METHODS_FORMAT_AUDIT_REPORT.md").write_text(
        "FORMAT_LOCALLY_READY\n",
        encoding="utf-8",
    )
    release_archives = root / "release" / "archives"
    release_archives.mkdir(parents=True)
    (release_archives / "sheafsignal_zenodo_upload.zip").write_bytes(b"zenodo")
    (release_archives / "sheafsignal_github_release.zip").write_bytes(b"github")
    figures = root / "figures"
    figures.mkdir()
    figure_rows = []
    for idx in range(1, 6):
        figure = figures / f"figure{idx}.pdf"
        figure.write_bytes(b"%PDF-1.4\n%%EOF\n")
        figure_rows.append(
            {
                "figure_id": f"fig{idx}",
                "path": figure.relative_to(root).as_posix(),
                "source_results_dir": "benchmarks/results",
            }
        )
    pd.DataFrame(figure_rows).to_csv(
        manuscript / "figure_manifest.tsv",
        sep="\t",
        index=False,
    )

    paths = script.build_package(
        root=root,
        output_dir=manuscript / "submission_upload_package",
        main_md=main_md,
        cover_md=cover_md,
        supplement_md=supplement_md,
    )

    for path in paths.values():
        assert path.exists()
    assert "Decision: `PASS`" in paths["text_check"].read_text(encoding="utf-8")
    layout = pd.read_csv(paths["layout_check"], sep="\t")
    assert set(layout["docx_key"]) == {"main_docx", "cover_docx", "supplement_docx"}
    figure_upload = pd.read_csv(paths["figure_manifest"], sep="\t")
    assert len(figure_upload) == 5
    assert set(figure_upload["status"]) == {"copied"}
    manifest = pd.read_csv(paths["manifest"], sep="\t")
    assert set(manifest["upload_item"]) >= {
        "main_manuscript",
        "cover_letter",
        "supplementary_information",
        "fig1",
    }
